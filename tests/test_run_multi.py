"""Tests for the multi-config driver's SLURM-facing behaviour."""

import json
import os
import re
import sys
import time
from pathlib import Path

sys.path.insert(
    0, str(Path(__file__).resolve().parents[1] / "workflow" / "scripts")
)

import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
import pytest  # noqa: E402
import yaml  # noqa: E402

from run_multi import (  # noqa: E402
    _CONVERT_KEYS,
    _relate_cmd,
    _snakemake_cmd,
    _validate_configs,
    slurm_jobname_prefix,
)

# The SLURM executor's own rule (snakemake_executor_plugin_slurm): it raises a
# WorkflowError and aborts the whole run if the prefix does not match.
_EXECUTOR_RULE = re.compile(r"^[A-Za-z0-9_-]{1,50}$")


def test_jobname_prefix_satisfies_the_executor():
    """Whatever a label_name contains, the prefix must stay submittable.

    The executor names jobs after a UUID and refuses a --job-name override, so
    this prefix is the only thing that makes squeue readable -- and an invalid
    one fails the run rather than degrading.
    """
    for label in ("nuclei_labels", "cyto_labels", "convert", "a"):
        assert _EXECUTOR_RULE.match(slurm_jobname_prefix(label))

    # Characters a label might plausibly pick up are sanitised, not passed on.
    assert _EXECUTOR_RULE.match(slurm_jobname_prefix("cilia/v2 (test)"))
    assert _EXECUTOR_RULE.match(slurm_jobname_prefix("run 1: nuclei"))
    # And an over-long label is truncated to the executor's 50-char limit.
    assert _EXECUTOR_RULE.match(slurm_jobname_prefix("x" * 200))


def test_jobname_prefix_keeps_the_label_readable():
    """The label must lead, since that is what a queue listing truncates to."""
    assert slurm_jobname_prefix("nuclei_labels") == "pw-nuclei_labels"
    assert slurm_jobname_prefix("convert") == "pw-convert"


def test_common_configfile_is_merged_under_the_per_config_one():
    """Snakemake merges --configfile values in order, later winning.

    That ordering is the whole mechanism: shared settings come from common.yaml
    and the per-config file overrides only what differs. Swap the two and every
    config would silently get the shared defaults instead of its own channel.
    """
    cmd = _snakemake_cmd(
        Path("config/config_nuclei.yaml"),
        workflow_dir=Path("workflow"),
        profile=None,
        cores=8,
        dry_run=False,
        common=Path("config/common.yaml"),
    )
    i = cmd.index("--configfile")
    assert cmd[i + 1].endswith("common.yaml")
    assert cmd[i + 2].endswith("config_nuclei.yaml")

    # Without a common file the invocation is unchanged: one configfile, so
    # a self-contained config keeps working exactly as before.
    plain = _snakemake_cmd(
        Path("config/config_nuclei.yaml"),
        workflow_dir=Path("workflow"),
        profile=None,
        cores=8,
        dry_run=False,
    )
    j = plain.index("--configfile")
    assert plain[j + 1].endswith("config_nuclei.yaml")
    assert not plain[j + 2].endswith(".yaml")


def test_convert_keys_must_agree_across_configs():
    """`convert` runs once from the first config, so a later one is ignored.

    Setting shard on the second config and watching a million files appear
    anyway is invisible without this check -- there is no log line saying the
    value was dropped, because nothing ever read it.
    """
    paths = [Path("a.yaml"), Path("b.yaml")]
    base = {"work_dir": "/w", "tile_shape": [16, 512, 512], "level": 0}
    good = [
        {**base, "label_name": "a", "shard": True},
        {**base, "label_name": "b", "shard": True},
    ]
    assert _validate_configs(paths, good) == "/w"

    bad = [
        {**base, "label_name": "a", "shard": True},
        {**base, "label_name": "b", "shard": False},
    ]
    # It reports every problem and exits, rather than raising, so that a
    # mistake costs one readable message instead of a traceback.
    with pytest.raises(SystemExit):
        _validate_configs(paths, bad)


def test_shipped_multi_configs_are_consistent():
    """The shipped example must satisfy its own validator.

    It is the thing users copy, so a config set that run_multi would refuse to
    start is worse than no example at all.
    """
    cfg_dir = Path(__file__).resolve().parents[1] / "workflow" / "config"
    multi = yaml.safe_load((cfg_dir / "multi.yaml").read_text())
    common = yaml.safe_load((cfg_dir.parent / multi["common"]).read_text())
    paths = [cfg_dir.parent / p for p in multi["segmentations"]]
    cfgs = [{**common, **yaml.safe_load(p.read_text())} for p in paths]

    assert _validate_configs(paths, cfgs) == common["work_dir"]
    # Every key convert reads comes from the shared file, not a per-config one.
    for path in paths:
        own = yaml.safe_load(path.read_text())
        assert not set(own) & set(_CONVERT_KEYS), path.name


def _workflow_dir() -> Path:
    return Path(__file__).resolve().parents[1] / "workflow"


def test_occupancy_is_a_submitted_rule_not_a_localrule():
    """The occupancy build must never run on the submit host.

    It streams the entire image. Doing that in the run_multi driver ran it on
    a login node, where a multi-terabyte read is killed with no traceback --
    the run just returned to the prompt. Only fetch_model may be local (it
    needs network); everything else has to get a real allocation.
    """
    wf = _workflow_dir()
    snakefile = (wf / "Snakefile").read_text()
    local_block = snakefile.split("localrules:")[1].split("rule ")[0]
    local = {
        line.strip().rstrip(",")
        for line in local_block.splitlines()
        if line.strip() and not line.strip().startswith("#")
    }
    assert local == {"fetch_model"}, local

    rules = (wf / "rules" / "convert.smk").read_text()
    assert "rule occupancy:" in rules


def test_occupancy_is_not_rebuilt_by_the_driver():
    """run_multi must ask Snakemake for the map, not build it in-process.

    An in-process build bypasses the scheduler entirely, which is how it ended
    up on the login node.
    """
    src = (_workflow_dir() / "scripts" / "run_multi.py").read_text()
    assert "build_occupancy_map(" not in src
    assert "occupancy.zarr" in src


def test_relate_is_submitted_via_slurm_under_profile():
    """The relate step must never run in-process on the submit host.

    Same failure mode as the occupancy map: label_relations() streams every
    chunk of two full-resolution label volumes. A prior fix moved the map
    build off the login node; the relate step made the identical mistake and
    hung there for the same reason until this fix.
    """
    src = (_workflow_dir() / "scripts" / "run_multi.py").read_text()
    assert "from relate import run_relations" in src
    assert '"srun"' in src
    assert "label_relations(" not in src


def _relate_kwargs(**overrides):
    kwargs = dict(
        work_dir="/w",
        image_store="/w/image.zarr",
        workflow_dir=Path("/workflow"),
        relate_partition="scicore",
        relate_mem="32G",
        relate_cpus=8,
        relate_time=180,
        relate_qos=None,
    )
    kwargs.update(overrides)
    return kwargs


def test_relate_cmd_is_one_job_per_relation_pair():
    """A killed shared job used to lose every relation still queued behind

    the one that was running -- one srun per pair means a slow or failing
    pair can no longer starve, or take down, its siblings' time budget.
    """
    relations = [
        {"a": "nuclei_labels", "b": "cyto_labels", "output": "n2c.xlsx"},
        {"a": "cilia_labels", "b": "cyto_labels", "output": "c2c.xlsx"},
    ]
    cmds = [_relate_cmd(rel, **_relate_kwargs()) for rel in relations]

    assert len(cmds) == 2
    for cmd, rel in zip(cmds, relations):
        assert cmd[0] == "srun"
        assert "--relations" in cmd
        # Each job's payload is *only* its own pair, not the whole list.
        payload = json.loads(cmd[cmd.index("--relations") + 1])
        assert payload == [rel]


def test_relate_cmd_job_name_identifies_the_pair():
    cmd = _relate_cmd(
        {"a": "cilia_labels", "b": "cyto_labels"}, **_relate_kwargs()
    )
    name = cmd[cmd.index("--job-name") + 1]
    assert _EXECUTOR_RULE.match(name)
    assert "cilia_labels" in name
    assert "cyto_labels" in name


def test_relate_cmd_gives_each_pair_its_own_log():
    """Concurrent per-pair jobs sharing one relate.log would interleave --

    each pair's --log must be a distinct file, or the whole point of
    splitting the log the way segment/<batch>.log already does is lost.
    """
    cmds = [
        _relate_cmd(rel, **_relate_kwargs())
        for rel in (
            {"a": "nuclei_labels", "b": "cyto_labels"},
            {"a": "cilia_labels", "b": "cyto_labels"},
        )
    ]
    logs = [cmd[cmd.index("--log") + 1] for cmd in cmds]
    assert len(set(logs)) == 2
    assert all(log.startswith("/w/logs/relate/") for log in logs)


def test_relate_cmd_omits_qos_by_default():
    cmd = _relate_cmd({"a": "a", "b": "b"}, **_relate_kwargs())
    assert "--qos" not in cmd


def test_relate_cmd_passes_qos_when_set():
    """A default QOS whose MaxWall is shorter than --relate-time is exactly

    what killed a real run (QOSMaxWallDurationPerJobLimit) -- --relate-qos
    lets a longer one be requested explicitly instead of guessed at.
    """
    cmd = _relate_cmd({"a": "a", "b": "b"}, **_relate_kwargs(relate_qos="1day"))
    assert cmd[cmd.index("--qos") + 1] == "1day"


def test_relate_script_has_the_real_bookkeeping():
    """relate.py must be the actual implementation, not a stub.

    Submitting the wrong (or a trimmed-down) script would silently produce a
    workbook missing the unmatched-label rows the docstring promises.
    """
    src = (_workflow_dir() / "scripts" / "relate.py").read_text()
    assert "def run_relations(" in src
    assert "label_relations" in src
    assert "openpyxl" in src


def test_view_script_loads_every_label_by_default():
    """view.py must not override labels=, or auto-load stops working.

    view_in_napari's labels=None default is what auto-loads every label
    group under <image>/labels/<name>/ as its own layer -- passing an
    explicit labels= here would silently drop that and show only one.
    """
    src = (_workflow_dir() / "scripts" / "view.py").read_text()
    assert "from patchworks.plugins.napari import view_in_napari" in src
    assert "labels=" not in src


def test_viewer_is_a_separate_opt_in_pixi_environment():
    """napari's Qt/GUI deps must stay out of the default headless env.

    Adding them to the default `[pypi-dependencies]` would pull heavy GUI
    dependencies into every SLURM job's environment for a feature only used
    interactively.
    """
    src = (_workflow_dir() / "pixi.toml").read_text()
    default_deps = src.split("[pypi-dependencies]")[1].split("[feature")[0]
    assert "napari" not in default_deps
    assert 'viewer = { features = ["viewer"] }' in src
    assert 'extras = ["napari"]' in src


def test_relate_writes_its_own_log():
    """relate.py runs via srun, not a Snakemake rule -- nothing else wires up

    its logging (see prepare/segment/merge's `log:` directives), so main()
    has to call start_log() itself or its output only ever streams to
    whatever invoked srun and is gone once that terminal scrolls past it.
    """
    src = (_workflow_dir() / "scripts" / "relate.py").read_text()
    assert "from _pw import start_log" in src
    assert "start_log(" in src
    assert '"logs" / "relate.log"' in src


def test_relate_rechunks_mismatched_label_arrays(tmp_path):
    """A chunk-layout mismatch must be rechunked away, not require a re-run.

    Two configs are free to have segmented at different tile_shape (one
    published before the other's config changed, or a cheaper method sized
    its own tile differently) -- label_relations() itself refuses mismatched
    chunks by design, but that only means the caller has to rechunk one side
    first, not that the whole segmentation needs redoing.
    """
    import zarr

    from relate import run_relations

    image_store = str(tmp_path / "image.zarr")

    # a: labels 1 and 2, split at x=5. b: a single label 10 covering all of
    # a's label 1 and none of label 2 -- built with a *different* chunking.
    a_data = np.zeros((1, 10), dtype=np.int32)
    a_data[0, :5] = 1
    a_data[0, 5:] = 2
    b_data = np.zeros((1, 10), dtype=np.int32)
    b_data[0, :5] = 10

    root = zarr.open_group(image_store, mode="w")
    labels = root.require_group("labels")
    a_grp = labels.require_group("nuclei_labels")
    a_arr = a_grp.create_array(
        name="0", shape=a_data.shape, chunks=(1, 2), dtype=np.int32
    )
    a_arr[:] = a_data
    a_grp.attrs["sequential_labels"] = True
    a_grp.attrs["n_objects"] = 2

    b_grp = labels.require_group("cyto_labels")
    b_arr = b_grp.create_array(
        name="0", shape=b_data.shape, chunks=(1, 5), dtype=np.int32
    )
    b_arr[:] = b_data
    b_grp.attrs["sequential_labels"] = True
    b_grp.attrs["n_objects"] = 1

    out_dir = tmp_path / "work"
    out_dir.mkdir()
    run_relations(
        str(out_dir),
        image_store,
        [{"a": "nuclei_labels", "b": "cyto_labels", "output": "rel.xlsx"}],
    )

    wb = openpyxl.load_workbook(out_dir / "rel.xlsx")
    rows = {
        row[0]: (row[1], row[2], row[3])
        for row in wb["nuclei_labels"].iter_rows(min_row=2, values_only=True)
    }
    assert rows[1] == (10, 5, 1.0)  # label 1 fully inside b's label 10
    assert rows[2] == (None, 0, 0)  # label 2 touches nothing in b


def test_relation_up_to_date_missing_output_is_false(tmp_path):
    from relate import _relation_up_to_date

    assert not _relation_up_to_date(
        str(tmp_path), "a", "b", tmp_path / "nope.xlsx"
    )


def test_relation_up_to_date_missing_marker_is_false(tmp_path):
    """No labels.done for a label means its state can't be judged -- treat

    that as "recompute", not as "trust the existing workbook".
    """
    from relate import _relation_up_to_date

    out = tmp_path / "rel.xlsx"
    out.write_text("x")
    assert not _relation_up_to_date(str(tmp_path), "a", "b", out)


def test_relation_up_to_date_true_only_when_newer_than_both_markers(
    tmp_path,
):
    from relate import _relation_up_to_date

    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    (tmp_path / "a" / "labels.done").touch()
    (tmp_path / "b" / "labels.done").touch()
    out = tmp_path / "rel.xlsx"
    out.write_text("x")

    now = time.time()
    os.utime(tmp_path / "a" / "labels.done", (now, now))
    os.utime(tmp_path / "b" / "labels.done", (now, now))

    # older than both markers -> stale
    os.utime(out, (now - 10, now - 10))
    assert not _relation_up_to_date(str(tmp_path), "a", "b", out)

    # newer than both markers -> up to date
    os.utime(out, (now + 10, now + 10))
    assert _relation_up_to_date(str(tmp_path), "a", "b", out)

    # b re-merged after the workbook was written -> stale again
    os.utime(tmp_path / "b" / "labels.done", (now + 20, now + 20))
    assert not _relation_up_to_date(str(tmp_path), "a", "b", out)


def test_relate_skips_a_relation_whose_workbook_is_up_to_date(tmp_path):
    """A retry must not recompute what already finished -- only what a

    shared, now-split-per-pair job left missing after a partial failure.
    Proven by planting sentinel content no real relation would produce: if
    run_relations() recomputed anyway, the sentinel would be gone.
    """
    from relate import run_relations

    work_dir = tmp_path / "work"
    work_dir.mkdir()
    for name in ("a_labels", "b_labels"):
        (work_dir / name).mkdir()
        (work_dir / name / "labels.done").touch()

    out_path = work_dir / "rel.xlsx"
    sentinel = openpyxl.Workbook()
    sentinel.active.append(["sentinel"])
    sentinel.save(out_path)
    os.utime(out_path, (time.time() + 60, time.time() + 60))

    # No image_store/labels at all -- if this weren't skipped, run_relations
    # would raise trying to open them, not just produce the wrong content.
    run_relations(
        str(work_dir),
        str(tmp_path / "image.zarr"),
        [{"a": "a_labels", "b": "b_labels", "output": "rel.xlsx"}],
    )

    wb = openpyxl.load_workbook(out_path)
    assert wb.active["A1"].value == "sentinel"


def test_mixed_nuclei_channel_auto_passes_validation():
    """A channel-count mismatch under `tile_shape: "auto"` is no longer

    refused at validation time -- it's resolved automatically instead (see
    `_resolve_shared_tile_shape`), which needs the converted image's real
    shape/dtype and so can only run after phase A, not from
    `_validate_configs()`. This used to `sys.exit` here; asserting that
    would now be testing the wrong layer.
    """
    paths = [Path("a.yaml"), Path("b.yaml")]
    base = {"work_dir": "/w", "tile_shape": "auto", "level": 0}

    mixed = [
        {**base, "label_name": "a", "channel": 0, "nuclei_channel": 1},
        {**base, "label_name": "b", "channel": 2},
    ]
    assert _validate_configs(paths, mixed) == "/w"

    # Same pair with one explicit shape is fine: both get that tile.
    pinned = [{**c, "tile_shape": [16, 512, 512]} for c in mixed]
    assert _validate_configs(paths, pinned) == "/w"

    # And "auto" is fine when every config carries the same channel count.
    both = [{**mixed[0]}, {**mixed[1], "nuclei_channel": 3}]
    assert _validate_configs(paths, both) == "/w"


def test_resolve_shared_tile_shape_pins_the_smallest_candidate(
    tmp_path, monkeypatch
):
    """The shared tile must be the tightest of every config's own budget.

    A larger tile than some config's own sizer output would ask that config
    for more memory than its settings were judged to need -- only the
    smallest candidate is safe for every config at once.
    """
    import run_multi

    class _FakeImage:
        shape = (10, 100, 100)
        dtype = "uint16"

    calls = []

    def _fake_load_ome_zarr(store, *, channel, level):
        calls.append((store, channel, level))
        return _FakeImage()

    # One tile per config, matched up by call order (channel 0 then 1).
    fake_tiles = [(8, 64, 64), (4, 32, 32)]

    def _fake_sizer_cellpose(shape, dtype, **kwargs):
        return fake_tiles[len(calls) - 1]

    monkeypatch.setattr(
        "patchworks.load_ome_zarr", _fake_load_ome_zarr, raising=False
    )
    monkeypatch.setattr(
        "patchworks.auto_tile_shape_cellpose",
        _fake_sizer_cellpose,
        raising=False,
    )

    seg_cfgs = [
        {
            "channel": 0,
            "level": 0,
            "method": "cellpose",
            "cellpose": {"do_3D": True, "gpu": True},
        },
        {
            "channel": 1,
            "level": 0,
            "nuclei_channel": 2,
            "method": "cellpose",
            "cellpose": {"do_3D": True, "gpu": True},
        },
    ]
    out = run_multi._resolve_shared_tile_shape(
        seg_cfgs, "/w/image.zarr", str(tmp_path)
    )
    assert out == tmp_path / ".multi_tile_shape.generated.yaml"
    written = yaml.safe_load(out.read_text())
    assert written == {"tile_shape": [4, 32, 32]}  # the smaller candidate
    assert len(calls) == 2


def test_merge_reshards_level_zero_when_asked():
    """`shard_labels` has to reshard level 0 *after* everything wrote to it.

    Ordering is the whole correctness argument: the merge and the volume
    filter both rewrite level 0 a chunk at a time, so a shard created before
    either of them would be read-modify-written by several writers and lose
    chunks. It also has to land before `register_labels`, so the pyramid is
    built from the level that will actually be on disk.
    """
    src = (_workflow_dir() / "scripts" / "merge.py").read_text()
    assert 'cfg.get("shard_labels", False)' in src
    assert "reshard_level(" in src

    reshard = src.index("reshard_level(label_group")
    assert src.index("merge_tile_labels(") < reshard
    assert src.index("filter_labels_by_size(") < reshard
    assert reshard < src.index("register_labels(\n")


def test_merge_shards_the_label_pyramid():
    """`shard:` has to reach the label pyramid, not just the conversion.

    Only the pyramid levels can take it: level 0 is written one chunk at a
    time by concurrent segment jobs (or the merge's own pool), and a shard
    must be written whole by a single writer -- see _write_pyramid's note.
    Levels 1..N go through one dask pass, so they can be sharded.
    """
    src = (_workflow_dir() / "scripts" / "merge.py").read_text()
    # The call's own arguments contain ")", so end it on the closing line.
    register = src.split("register_labels(")[1].split("\n)")[0]
    assert 'shard=cfg.get("shard", False)' in register


def test_ngff_version_reaches_convert_and_merge():
    """Both writers need it, or a store ends up half one version.

    convert makes the image; merge adds the label pyramid to it, in a
    separate job hours later. If only one of them read the key the store
    would mix zarr v2 arrays with v3 ones.
    """
    convert = (_workflow_dir() / "scripts" / "convert.py").read_text()
    merge = (_workflow_dir() / "scripts" / "merge.py").read_text()
    for name, src in (("convert.py", convert), ("merge.py", merge)):
        assert 'ngff_version=cfg.get("ngff_version", "auto")' in src, name


def test_ngff_version_must_agree_across_configs():
    """It decides the store's format, and convert runs once, from config #1."""
    from run_multi import _CONVERT_KEYS

    assert "ngff_version" in _CONVERT_KEYS


def test_store_marker_file_follows_the_zarr_format():
    """The rules watch a file whose *name* is the zarr format's.

    zarr v3 writes "zarr.json", v2 writes ".zgroup". NGFF 0.4 means a v2
    store, so a hardcoded zarr.json would leave `convert` waiting forever on
    a file that is never written.
    """
    smk = (_workflow_dir() / "rules" / "common.smk").read_text()
    assert "ZARR_ROOT_FILE" in smk
    assert '".zgroup"' in smk
    assert '"zarr.json"' in smk
    assert 'IMAGE_OK = f"{IMAGE}/{ZARR_ROOT_FILE}"' in smk
    assert 'OCCUPANCY_OK = f"{OCCUPANCY}/{ZARR_ROOT_FILE}"' in smk

    # Both scripts that turn the marker back into a store path must strip
    # whichever name is in use.
    for script in ("convert.py", "build_occupancy.py"):
        src = (_workflow_dir() / "scripts" / script).read_text()
        assert '.removesuffix("/zarr.json")' in src, script
        assert '.removesuffix("/.zgroup")' in src, script
